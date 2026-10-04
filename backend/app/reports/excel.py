from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import get_settings
from app.core.timezone import now_local
from app.models.stoppage import StoppageRecord


HEADER_FILL = PatternFill("solid", fgColor="1B3A4B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1B3A4B")
THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
TOTAL_FILL = PatternFill("solid", fgColor="EEF2F6")


def _header(ws, headers: list[str], row: int = 1) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"
    ws.freeze_panes = f"A{row + 1}"
    ws.row_dimensions[row].height = 22


def _autosize(ws) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = 12
        for cell in column:
            if cell.value:
                width = max(width, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[letter].width = width


def _meta(ws, title: str, filters_text: str) -> int:
    settings = get_settings()
    ws["A1"] = settings.app_name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = f"Company: {settings.company_name}"
    ws["A4"] = f"Generated: {now_local().strftime('%d %b %Y %I:%M %p')} (Asia/Kolkata)"
    ws["A5"] = f"Filters: {filters_text}"
    return 7


def format_filters(records_meta: dict) -> str:
    parts = []
    for key, value in records_meta.items():
        if value:
            parts.append(f"{key}={value}")
    return ", ".join(parts) if parts else "None"


def write_raw_sheet(ws, records: list[StoppageRecord], start_row: int = 1) -> None:
    headers = [
        "Production Date",
        "Shift",
        "Supervisor",
        "Machine",
        "Stoppage Reason",
        "Duration (Minutes)",
        "Details",
        "Remarks",
        "Created At",
    ]
    _header(ws, headers, start_row)
    total = 0
    for offset, record in enumerate(records, 1):
        row = start_row + offset
        created = now_local(record.created_at)
        values = [
            record.production_date,
            record.shift.code,
            record.supervisor.name,
            record.machine.code,
            record.reason.name,
            float(record.duration_minutes),
            record.details or "",
            record.remarks or "",
            created.replace(tzinfo=None),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = THIN
            if col in (7, 8):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=1).number_format = "DD MMM YYYY"
        ws.cell(row=row, column=6).number_format = "0"
        ws.cell(row=row, column=9).number_format = "DD MMM YYYY HH:MM AM/PM"
        total += float(record.duration_minutes)
    total_row = start_row + len(records) + 1
    ws.cell(row=total_row, column=5, value="Total").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=6, value=total)
    total_cell.font = Font(bold=True)
    total_cell.number_format = "0"
    for col in range(1, 10):
        ws.cell(row=total_row, column=col).fill = TOTAL_FILL
        ws.cell(row=total_row, column=col).border = THIN
    _autosize(ws)


def write_named_summary(ws, title: str, rows: list[tuple], filters_text: str) -> None:
    start = _meta(ws, title, filters_text)
    headers = ["Name", "Total Downtime (Minutes)", "Stoppage Count"]
    _header(ws, headers, start)
    total_down = 0
    total_count = 0
    for offset, (name, downtime, count) in enumerate(rows, 1):
        row = start + offset
        ws.cell(row=row, column=1, value=name).border = THIN
        cell = ws.cell(row=row, column=2, value=float(downtime))
        cell.border = THIN
        cell.number_format = "0"
        ws.cell(row=row, column=3, value=int(count)).border = THIN
        total_down += float(downtime)
        total_count += int(count)
    total_row = start + len(rows) + 1
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=total_down).font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=total_count).font = Font(bold=True)
    for col in range(1, 4):
        ws.cell(row=total_row, column=col).fill = TOTAL_FILL
        ws.cell(row=total_row, column=col).border = THIN
    _autosize(ws)


def build_workbook(report_type: str, records: list[StoppageRecord], summaries: dict, filters_text: str) -> bytes:
    wb = Workbook()
    if report_type in ("raw", "complete"):
        raw = wb.active
        raw.title = "Raw Data"
        start = _meta(raw, "Stoppage Raw Data", filters_text)
        write_raw_sheet(raw, records, start)
    else:
        wb.remove(wb.active)

    if report_type in ("summary", "complete"):
        sheets = [
            ("Machine Summary", summaries["machines"]),
            ("Reason Summary", summaries["reasons"]),
            ("Shift Summary", summaries["shifts"]),
            ("Daily Summary", summaries["daily"]),
        ]
        for title, rows in sheets:
            ws = wb.create_sheet(title)
            write_named_summary(ws, title, rows, filters_text)

    if report_type == "complete" and wb.worksheets:
        wb._sheets = wb._sheets  # keep default order
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
